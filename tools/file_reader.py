"""
file_reader.py — Read content from previously generated output files.

Supports .xlsx (returns all sheet data as text tables),
         .docx (returns all paragraph and table text),
         .pptx (returns slide titles and bullet text).
"""

import json
from pathlib import Path
from config import OUTPUT_DIR


def read_output_file(filename: str, output_dir=None) -> str:
    """
    Read the content of a file that was previously created in the output folder.

    Args:
        filename: The filename (with extension) of the file to read.
        output_dir: The caller's own output folder. Scoping reads to it keeps one
                    user from reading (or being offered) another user's documents.

    Returns:
        JSON string with the file content as structured text, or an error message.
    """
    base = (Path(output_dir) if output_dir else OUTPUT_DIR).resolve()
    safe_name = Path(filename).name
    filepath = (base / safe_name).resolve()

    # Security: only allow reading from the caller's own output directory
    if not str(filepath).startswith(str(base)):
        return json.dumps({"error": "Access denied: file is outside the output directory."})

    if not filepath.exists():
        available = [f.name for f in base.iterdir() if f.is_file()] if base.exists() else []
        return json.dumps({
            "error": f"File '{safe_name}' not found in output directory.",
            "available_files": available,
        })

    suffix = filepath.suffix.lower()

    try:
        if suffix == ".xlsx":
            return _read_excel(filepath)
        elif suffix == ".docx":
            return _read_word(filepath)
        elif suffix == ".pptx":
            return _read_pptx(filepath)
        else:
            return json.dumps({"error": f"Unsupported file type: {suffix}. Supported: .xlsx, .docx, .pptx"})
    except Exception as e:
        return json.dumps({"error": f"Failed to read file: {str(e)}"})


def _read_excel(filepath: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(filepath, data_only=True)
    sheets_out = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            # Skip completely empty rows
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                rows.append([str(c) if c is not None else "" for c in row])

        sheets_out.append({
            "sheet": sheet_name,
            "rows": rows,
        })

    return json.dumps({"type": "excel", "filename": filepath.name, "sheets": sheets_out}, indent=2)


def _read_word(filepath: Path) -> str:
    from docx import Document
    doc = Document(str(filepath))
    sections = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else "Normal"
            sections.append({"style": style, "text": text})

    tables_out = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables_out.append({"table_index": i, "rows": rows})

    return json.dumps({
        "type": "word",
        "filename": filepath.name,
        "paragraphs": sections,
        "tables": tables_out,
    }, indent=2)


def _read_pptx(filepath: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(filepath))
    slides_out = []

    for i, slide in enumerate(prs.slides):
        title = ""
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        if shape.shape_type == 13 or (hasattr(shape, "placeholder_format")
                                and shape.placeholder_format
                                and shape.placeholder_format.idx == 0):
                            title = t
                        else:
                            texts.append(t)
        slides_out.append({"slide": i + 1, "title": title, "content": texts})

    return json.dumps({"type": "powerpoint", "filename": filepath.name, "slides": slides_out}, indent=2)
