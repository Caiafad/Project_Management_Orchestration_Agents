"""
Validator tests against real .docx / .xlsx / .pptx files written to a temp dir.

Each test builds a document with one specific defect — the kind the agent
prompts call a failure — and asserts the validator reports it.
"""

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from tools.doc_validator import (cross_check_totals, validate_docx, validate_outputs,
                                 validate_pptx, validate_xlsx)

SCOPE_HEADINGS = ["Introduction", "Project Overview", "In-Scope", "Out-of-Scope",
                  "Functional Requirements", "Non-Functional Requirements",
                  "Assumptions and Constraints", "Acceptance Criteria",
                  "Scope Change Management"]


def write_docx(path, sections, tables=None):
    """sections: list of (heading, body). tables: {heading: [[row], ...]}"""
    doc = Document()
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        if body:
            doc.add_paragraph(body)
        for rows in (tables or {}).get(heading, []):
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = str(value)
    doc.save(str(path))
    return path


def write_xlsx(path, sheets):
    """sheets: {name: [[header...], [row...]]}"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return path


def write_pptx(path, slides):
    """slides: list of (title, [bullets])"""
    prs = Presentation()
    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        if bullets:
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(4))
            frame = box.text_frame
            frame.text = bullets[0]
            for bullet in bullets[1:]:
                frame.add_paragraph().text = bullet
    prs.save(str(path))
    return path


# ── Word ──────────────────────────────────────────────────────────────────────

def test_complete_scope_document_passes(tmp_path):
    body = " ".join(f"FR-{i:03d} The system SHALL do thing {i}." for i in range(1, 13))
    nfrs = " ".join(f"NFR-{i:03d} Performance target {i}." for i in range(1, 9))
    sections = [(h, "Substantive content for this section.") for h in SCOPE_HEADINGS]
    sections[4] = ("Functional Requirements", body)
    sections[5] = ("Non-Functional Requirements", nfrs)
    path = write_docx(tmp_path / "scope.docx", sections)

    result = validate_docx(path, {"required_headings": SCOPE_HEADINGS, "min_sections": 9,
                                  "min_numbered": [(r"\bFR-\d+", 8, "FRs"),
                                                   (r"\bNFR-\d+", 7, "NFRs")]})
    assert result.ok, result.issues


def test_missing_required_section_is_reported(tmp_path):
    kept = [h for h in SCOPE_HEADINGS if h != "Non-Functional Requirements"]
    path = write_docx(tmp_path / "scope.docx", [(h, "Content here.") for h in kept])

    result = validate_docx(path, {"required_headings": SCOPE_HEADINGS})
    assert not result.ok
    assert any("Non-Functional Requirements" in i for i in result.issues)


def test_heading_with_no_content_is_reported(tmp_path):
    path = write_docx(tmp_path / "doc.docx",
                      [("Introduction", "Real content."), ("Acceptance Criteria", "")])
    result = validate_docx(path)
    assert not result.ok
    assert any("Acceptance Criteria" in i and "no content" in i for i in result.issues)


@pytest.mark.parametrize("placeholder", ["TBD", "[insert cost here]", "Lorem ipsum dolor", "TODO"])
def test_placeholder_text_is_reported(tmp_path, placeholder):
    path = write_docx(tmp_path / "doc.docx", [("Budget", f"The total is {placeholder}.")])
    result = validate_docx(path)
    assert not result.ok
    assert any("placeholder" in i.lower() for i in result.issues)


def test_nfr_numbering_does_not_satisfy_the_fr_minimum(tmp_path):
    # NFR-001 contains the substring "FR-001"; the \b anchor must exclude it.
    nfrs = " ".join(f"NFR-{i:03d} requirement." for i in range(1, 13))
    path = write_docx(tmp_path / "doc.docx", [("Non-Functional Requirements", nfrs)])
    result = validate_docx(path, {"min_numbered": [(r"\bFR-\d+", 8, "functional requirements")]})
    assert not result.ok
    assert any("functional requirements" in i for i in result.issues)


def test_header_only_table_is_reported(tmp_path):
    path = write_docx(tmp_path / "doc.docx", [("Costs", "Narrative.")],
                      tables={"Costs": [[["Role", "Cost"]]]})
    result = validate_docx(path)
    assert not result.ok
    assert any("no data rows" in i for i in result.issues)


def test_section_missing_its_required_table_is_reported(tmp_path):
    path = write_docx(tmp_path / "doc.docx", [("Executive Summary", "Narrative only.")])
    result = validate_docx(path, {"required_tables": ["Executive Summary"]})
    assert not result.ok
    assert any("must contain a table" in i for i in result.issues)


def test_unopenable_file_is_reported(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a real docx")
    assert not validate_docx(bad).ok


# ── Excel ─────────────────────────────────────────────────────────────────────

FINANCIAL_SHEETS = ["Summary", "Labor Detail", "Infrastructure", "Monthly Cash Flow", "Assumptions"]


def financial_workbook(tmp_path, *, with_labor_total=True):
    labor = [["Role", "Headcount", "Total Cost"], ["Engineer", 4, 400000]]
    if with_labor_total:
        labor.append(["TOTAL LABOR", "", 400000])
    return write_xlsx(tmp_path / "fin.xlsx", {
        "Summary": [["Category", "Most Likely"], ["Labor", 400000], ["GRAND TOTAL", 460000]],
        "Labor Detail": labor,
        "Infrastructure": [["Item", "Cost"], ["Cloud", 60000], ["TOTAL INFRASTRUCTURE", 60000]],
        "Monthly Cash Flow": [["Category", "Month 1", "TOTAL"], ["Labor", 40000, 400000],
                              ["TOTAL MONTHLY", 46000, 460000]],
        "Assumptions": [["Item", "Value"], ["Contingency", "15%"]],
    })


def test_complete_financial_workbook_passes(tmp_path):
    path = financial_workbook(tmp_path)
    result = validate_xlsx(path, {"required_sheets": FINANCIAL_SHEETS,
                                  "total_rows": {"Labor Detail": "TOTAL LABOR",
                                                 "Infrastructure": "TOTAL INFRASTRUCTURE"}})
    assert result.ok, result.issues


def test_missing_sheet_is_reported(tmp_path):
    path = write_xlsx(tmp_path / "fin.xlsx", {"Summary": [["A"], [1]]})
    result = validate_xlsx(path, {"required_sheets": FINANCIAL_SHEETS})
    assert not result.ok
    assert any("Labor Detail" in i for i in result.issues)


def test_missing_total_row_is_reported(tmp_path):
    path = financial_workbook(tmp_path, with_labor_total=False)
    result = validate_xlsx(path, {"total_rows": {"Labor Detail": "TOTAL LABOR"}})
    assert not result.ok
    assert any("TOTAL LABOR" in i for i in result.issues)


def test_sheet_with_only_headers_is_reported(tmp_path):
    path = write_xlsx(tmp_path / "fin.xlsx", {"Labor Detail": [["Role", "Cost"]]})
    result = validate_xlsx(path)
    assert not result.ok
    assert any("no data rows" in i for i in result.issues)


def test_placeholder_in_workbook_is_reported(tmp_path):
    path = write_xlsx(tmp_path / "fin.xlsx",
                      {"Summary": [["Category", "Cost"], ["Labor", "TBD"]]})
    result = validate_xlsx(path)
    assert not result.ok
    assert any("placeholder" in i.lower() for i in result.issues)


# ── PowerPoint ────────────────────────────────────────────────────────────────

def test_deck_within_slide_range_passes(tmp_path):
    slides = [(f"Slide {i}", [f"Point {j}" for j in range(3)]) for i in range(9)]
    path = write_pptx(tmp_path / "deck.pptx", slides)
    result = validate_pptx(path, {"min_slides": 8, "max_slides": 14})
    assert result.ok, result.issues


def test_too_few_slides_is_reported(tmp_path):
    path = write_pptx(tmp_path / "deck.pptx", [("Only slide", ["A point"])])
    result = validate_pptx(path, {"min_slides": 8})
    assert not result.ok
    assert any("at least 8" in i for i in result.issues)


def test_too_many_bullets_is_reported(tmp_path):
    path = write_pptx(tmp_path / "deck.pptx",
                      [("Dense slide", [f"Bullet {i}" for i in range(9)])])
    result = validate_pptx(path, {"max_bullets_per_slide": 5})
    assert not result.ok
    assert any("bullet" in i.lower() for i in result.issues)


# ── Cross-document agreement ──────────────────────────────────────────────────

def test_matching_figures_pass_cross_check(tmp_path):
    xlsx = financial_workbook(tmp_path)
    docx = write_docx(tmp_path / "fin.docx", [("Executive Summary", "See table.")],
                      tables={"Executive Summary": [[["Scenario", "Total"],
                                                     ["Most Likely", "$460,000"]]]})
    assert cross_check_totals(xlsx, docx).ok


def test_figure_absent_from_the_workbook_is_reported(tmp_path):
    xlsx = financial_workbook(tmp_path)
    docx = write_docx(tmp_path / "fin.docx", [("Executive Summary", "See table.")],
                      tables={"Executive Summary": [[["Scenario", "Total"],
                                                     ["Most Likely", "$899,500"]]]})
    result = cross_check_totals(xlsx, docx)
    assert not result.ok
    assert any("899,500" in i for i in result.issues)


def test_validate_outputs_aggregates_by_filename(tmp_path):
    xlsx = write_xlsx(tmp_path / "fin.xlsx", {"Summary": [["Category", "Cost"]]})
    docx = write_docx(tmp_path / "fin.docx", [("Executive Summary", "")])
    report = validate_outputs([xlsx, docx], {"xlsx": {}, "docx": {}})
    assert not report["ok"]
    assert set(report["issues"]) == {"fin.xlsx", "fin.docx"}


def test_validator_never_raises_on_an_unknown_extension(tmp_path):
    other = tmp_path / "notes.txt"
    other.write_text("hello")
    assert validate_outputs([other], {})["ok"]
